"""
競合アカウント調査データを Excel ファイルに出力するスクリプト
"""

import datetime
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)

TODAY = datetime.date.today().strftime("%Y/%m/%d")

HEADER = [
    "発見日", "評価", "アカウント名", "URL", "FW",
    "開始日", "リプ数", "サブ/URL", "LINE",
    "note/無料", "note/有料", "note/アカウント", "ペルソナ", "コメント"
]

ACCOUNTS = [
    # ── Web 調査で発掘できたアカウント ────────────────────────────────────
    [TODAY, "B", "みつい",
     "https://x.com/mrmitsuikun", "2000", "2023/01", "50〜80", "",
     "なし", "あり", "なし", "https://note.com/mrmitsui",
     "AI美女（架空モデル）",
     "AI美女×X運用の記録をnoteで公開。2000FW達成実績あり。2025年開始基準は要確認"],

    [TODAY, "C", "九鬼若芽",
     "https://x.com/kukiwakame_ai", "不明", "2022/03", "〜50", "",
     "なし", "あり", "なし", "https://note.com/kukiwakame_ai",
     "AI美少女白書（イラスト系）",
     "AI美少女イラスト特化。2022年開始のため2025年条件外の可能性あり要確認"],

    [TODAY, "C", "怠目藍成（あいなる）",
     "https://x.com/darume_ainaru", "不明", "2025/01以降", "〜30", "",
     "要確認", "あり", "なし", "https://note.com/darume_ainaru",
     "AI美女研究・勉強中",
     "AI美女イラスト・X運用に関するnote多数。LINE有無は要確認"],

    [TODAY, "C", "てんねん",
     "https://x.com/munou_ac", "41000", "2024/頃", "〜100", "",
     "なし", "あり", "あり", "https://note.com/munou_ac",
     "AIイラストクリエイター",
     "AIイラスト1日12投稿。FW多すぎで対象外。参考競合として記録"],

    [TODAY, "C", "アミラナ",
     "https://x.com/amirana39", "不明", "2025/頃", "〜30", "",
     "要確認", "あり", "なし", "https://note.com/amirana39",
     "AIキャラ・日常系",
     "X AIコンパニオン機能紹介など。note有。詳細はX/Grokで要確認"],

    # ── Grok で追加調査が必要なプレースホルダー (6〜20) ────────────────────
    *[
        [TODAY, "未調査", f"（未調査 {i}）", "", "", "2025/", "", "",
         "要確認", "要確認", "要確認", "",
         "",
         'Grokで「日常系女子 note LINE 2025」「AIキャラ 架空ペルソナ note LINE」などで発掘']
        for i in range(6, 21)
    ],
]


def make_thin_border():
    thin = Side(style="thin")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def build_excel(output_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "競合調査"

    # ── ヘッダー行 ──────────────────────────────────────────────────────
    header_fill   = PatternFill("solid", fgColor="1F4E79")
    header_font   = Font(bold=True, color="FFFFFF", size=10)
    center_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align    = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill   = header_fill
        cell.font   = header_font
        cell.alignment = center_align
        cell.border = make_thin_border()

    # ── データ行 ──────────────────────────────────────────────────────
    eval_colors = {
        "A": "C6EFCE",   # 緑
        "B": "FFEB9C",   # 黄
        "C": "FCE4D6",   # 薄オレンジ
        "未調査": "F2F2F2",  # グレー
    }

    for row_idx, row in enumerate(ACCOUNTS, start=2):
        eval_val = row[1] if len(row) > 1 else ""
        fill_color = eval_colors.get(eval_val, "FFFFFF")
        row_fill = PatternFill("solid", fgColor=fill_color)

        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill      = row_fill
            cell.alignment = left_align
            cell.border    = make_thin_border()
            cell.font      = Font(size=9)

    # ── 列幅の調整 ──────────────────────────────────────────────────────
    col_widths = [12, 8, 18, 36, 8, 10, 8, 36, 8, 10, 10, 36, 22, 50]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[
            openpyxl.utils.get_column_letter(col_idx)
        ].width = width

    ws.row_dimensions[1].height = 28

    # ── ウィンドウ枠固定（1行目を固定）──────────────────────────────────
    ws.freeze_panes = "A2"

    wb.save(output_path)
    print(f"✅ Excel ファイルを作成しました: {output_path}")
    print(f"   - 行数: {len(ACCOUNTS)} 件（調査済み5件 + 未調査プレースホルダー15件）")


if __name__ == "__main__":
    output = "competitor_research.xlsx"
    build_excel(output)
